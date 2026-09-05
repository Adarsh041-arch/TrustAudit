"""DOCX and PDF report builders over enriched audit payloads (ported from V1 report_builders).

V1 bug not reproduced: `logger` is defined here so the reportlab-missing
fallback path cannot raise NameError.
"""
from __future__ import annotations

import logging
from datetime import UTC, datetime
from io import BytesIO
from typing import Any

logger = logging.getLogger(__name__)

try:
    from reportlab.lib import colors  # type: ignore[import-untyped]
    from reportlab.lib.pagesizes import letter  # type: ignore[import-untyped]
    from reportlab.lib.styles import ParagraphStyle  # type: ignore[import-untyped]
    from reportlab.platypus import (  # type: ignore[import-untyped]
        Paragraph,
        SimpleDocTemplate,
        Table,
        TableStyle,
    )

    REPORTLAB_AVAILABLE = True
except ImportError:  # pragma: no cover
    REPORTLAB_AVAILABLE = False

SEVERITY_COLORS = {
    "critical": "#993C1D", "high": "#C2410C", "medium": "#D97706", "low": "#6B7280",
}


def _per_doc_rows(doc: dict[str, Any]) -> list[list[str]]:
    rows = []
    for r in doc.get("failed_rules", []):
        rows.append([
            r.get("rule_id", ""), r.get("severity", ""),
            f"{r.get('finding', '')} — evidence: {r.get('evidence', '')}",
            r.get("impact", "") or r.get("rule_title", ""),
        ])
    return rows


def _score_label(document: dict[str, Any]) -> str:
    value = document.get("score")
    return "Not audited" if value is None else f"{float(value):.1f}%"


def generate_docx_report(payload: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, RGBColor

    docs = payload.get("documents", [])
    doc = Document()
    doc.add_heading(payload.get("audit_title", "Audit V2 Report"), level=0)
    stamp = datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")
    doc.add_paragraph(f"Generated: {stamp} — {len(docs)} document(s)")
    scored = [float(d["score"]) for d in docs if d.get("score") is not None]
    avg = sum(scored) / len(scored) if scored else None
    doc.add_paragraph(
        f"Overall score: {avg:.1f}%" if avg is not None else "Overall score: Not audited"
    )
    if payload.get("executive_summary"):
        doc.add_heading("Executive summary", level=1)
        doc.add_paragraph(str(payload["executive_summary"]))

    for d in docs:
        doc.add_heading(d.get("document_name", "?"), level=1)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(
            f"Score: {_score_label(d)} — {d.get('risk_level', '')} — "
            f"passed: {d.get('passed')}"
        )
        run.bold = True
        if d.get("summary_text"):
            doc.add_paragraph(str(d["summary_text"]))
        ml = d.get("ml_prediction") or {}
        doc.add_paragraph(f"Prediction: {ml.get('prediction', 'N/A')} ({ml.get('mode', '')})")
        if d.get("human_review_recommended"):
            doc.add_paragraph("HUMAN REVIEW RECOMMENDED")
        rows = _per_doc_rows(d)
        if rows:
            table = doc.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            for i, title in enumerate(["Rule", "Severity", "Finding & Evidence", "Impact"]):
                hdr[i].text = title
            for row in rows:
                cells = table.add_row().cells
                for i, cell in enumerate(row):
                    cells[i].text = cell
                sev_color = SEVERITY_COLORS.get(cells[1].text.lower(), "#6B7280")
                sev_runs = cells[1].paragraphs[0].runs
                if sev_runs:
                    sev_runs[0].font.color.rgb = RGBColor.from_string(sev_color.lstrip("#"))
            if d.get("preview_base64"):
                try:
                    import base64
                    buf = BytesIO(base64.b64decode(d["preview_base64"]))
                    doc.add_picture(buf, width=Inches(1.5))
                except Exception:
                    pass

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_pdf_report(payload: dict[str, Any]) -> bytes:
    if not REPORTLAB_AVAILABLE:  # pragma: no cover
        logger.warning("reportlab not installed; falling back to DOCX report")
        return generate_docx_report(payload)
    docs = payload.get("documents", [])
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=letter, leftMargin=54, rightMargin=54, topMargin=54, bottomMargin=54
    )
    title = ParagraphStyle(
        "Title", fontName="Helvetica-Bold", fontSize=16, leading=20,
        textColor=colors.HexColor("#0F6E56"),
    )
    body = ParagraphStyle("Body", fontName="Helvetica", fontSize=9, leading=12)
    story = [Paragraph(payload.get("audit_title", "Audit V2 Report"), title)]
    story.append(
        Paragraph(
            f"Generated: {datetime.now(UTC).strftime('%Y-%m-%d %H:%M UTC')} — "
            f"{len(docs)} document(s)",
            body,
        )
    )
    if payload.get("executive_summary"):
        story.append(Paragraph("Executive summary", ParagraphStyle(
            "Executive", parent=body, fontName="Helvetica-Bold", fontSize=11,
        )))
        story.append(Paragraph(str(payload["executive_summary"]), body))
    for d in docs:
        story.append(
            Paragraph(
                d.get("document_name", "?"),
                ParagraphStyle("Doc", parent=body, fontName="Helvetica-Bold", fontSize=11),
            )
        )
        if d.get("summary_text"):
            story.append(Paragraph(str(d["summary_text"]), body))
        story.append(
            Paragraph(
                f"Score: {_score_label(d)} — {d.get('risk_level', '')} — "
                f"passed: {d.get('passed')}",
                body,
            )
        )
        ml = d.get("ml_prediction") or {}
        story.append(
            Paragraph(
                f"Prediction: {ml.get('prediction', 'N/A')} ({ml.get('mode', '')})", body
            )
        )
        if d.get("human_review_recommended"):
            story.append(
                Paragraph(
                    "HUMAN REVIEW RECOMMENDED",
                    ParagraphStyle("HR", parent=body, textColor=colors.HexColor("#993C1D")),
                )
            )
        rows = [["Rule", "Severity", "Finding & Evidence", "Impact"]] + _per_doc_rows(d)
        if len(rows) > 1:
            for data_row in rows[1:]:
                sev = data_row[1]
                data_row[1] = Paragraph(
                    sev,
                    ParagraphStyle(
                        f"Sev-{sev or 'unknown'}",
                        parent=body,
                        textColor=colors.HexColor(SEVERITY_COLORS.get(sev.lower(), "#6B7280")),
                    ),
                )
            t = Table(rows, colWidths=[80, 60, 220, 144])
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5E7EB")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
            ]))
            story.append(t)
    doc.build(story)
    return buf.getvalue()
